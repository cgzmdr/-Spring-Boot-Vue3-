# -*- coding: utf-8 -*-
"""构建论文 DOCX：样式对齐模板（A4、宋体小四正文、黑体标题、固定行距 20 磅）。"""
import os
import sys
import importlib.util

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

ROOT = r'C:\codeDev\56\app'
OUT = os.environ.get(
    'THESIS_OUT',
    os.path.join(ROOT, '基于SpringBoot+Vue的56个民族文化数字化展示平台的设计与实现.docx'))

CONTENT_MODULES = ['content_ch1_3', 'content_ch3', 'content_ch4', 'content_ch5', 'content_ch6']


def load_blocks():
    blocks = []
    for name in CONTENT_MODULES:
        path = os.path.join(ROOT, '_thesis_build', 'tools', name + '.py')
        spec = importlib.util.spec_from_file_location(name, path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        blocks.extend(mod.BLOCKS)
    return blocks


# --------------------------------------------------------------------------- 工具函数
def set_run_font(run, ascii_font='Times New Roman', ea_font='宋体', size=None, bold=None, color=None):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), ea_font)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_outline_level(paragraph, level):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement('w:outlineLvl')
    el.set(qn('w:val'), str(level))
    pPr.append(el)


def shade(paragraph, fill='F5F6F8'):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def add_field(paragraph, instr, cached=''):
    """在段落中插入一个域（如 PAGE / TOC）。"""
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    fld_begin.set(qn('w:dirty'), 'true')
    run._r.append(fld_begin)

    run2 = paragraph.add_run()
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = instr
    run2._r.append(instr_el)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fld_sep)

    run4 = paragraph.add_run(cached)

    run5 = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run5._r.append(fld_end)
    return [run, run2, run3, run4, run5]


# --------------------------------------------------------------------------- 样式初始化
def init_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    pf = normal.paragraph_format
    pf.line_spacing = Pt(20)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specs = {
        'Heading 1': (22, 17, 16.5, 2.4),
        'Heading 2': (16, 13, 13, 1.73),
        'Heading 3': (14, 10, 10, 1.5),
    }
    for name, (size, sb, sa, ls) in specs.items():
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), '黑体')
        p = st.paragraph_format
        p.space_before = Pt(sb)
        p.space_after = Pt(sa)
        p.line_spacing = ls
        p.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def init_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)


def add_page_number_footer(section, start_at=None):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_run_font(r, size=Pt(10.5))
    for el in add_field(p, ' PAGE ', '1'):
        set_run_font(el, size=Pt(10.5))
    if start_at is not None:
        sectPr = section._sectPr
        pg = OxmlElement('w:pgNumType')
        pg.set(qn('w:start'), str(start_at))
        sectPr.append(pg)


# --------------------------------------------------------------------------- 内容写入
def add_body_paragraph(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    r = p.add_run(text)
    set_run_font(r, size=Pt(12))
    return p


def add_figure(doc, rel_path, caption, max_w_cm=15.0, max_h_cm=19.0):
    path = os.path.join(ROOT, rel_path.replace('/', os.sep))
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 关键：图片所在段落不能用「固定行距」，否则行框被压到 20 磅，
    # 内联图片会向上溢出页面而被裁掉。这里必须用单倍行距让行框随图片高度增长。
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(path, width=Cm(max_w_cm))
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_together = True
    cr = cap.add_run(caption)
    set_run_font(cr, size=Pt(10.5))
    return p


def add_table(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(2)
    cr = cap.add_run(caption)
    set_run_font(cr, size=Pt(10.5), bold=True)

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = Pt(14)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        r = p.add_run(h)
        set_run_font(r, size=Pt(9), bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            cell.text = ''
            lines = str(val).split('\n')
            for j, line in enumerate(lines):
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                p.paragraph_format.line_spacing = Pt(14)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.space_after = Pt(0)
                # 数字类短字段居中，名称与注释左对齐
                if i in (0, 2, 3, 4, 5) and len(str(line)) <= 12:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(line)
                set_run_font(r, size=Pt(9))
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
    return t


def add_code_block(doc, code):
    lines = code.split('\n')
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.left_indent = Pt(6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    shade(p, 'F4F5F7')
    for i, line in enumerate(lines):
        r = p.add_run(line.replace('\t', '    '))
        set_run_font(r, ascii_font='Consolas', ea_font='宋体', size=Pt(8.5))
        if i < len(lines) - 1:
            r.add_break()
    return p


def add_toc(doc, entries, page_map=None):
    """插入带缓存结果的 TOC 域：先用静态条目占位（Word 打开时自动更新为正式目录）。"""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

    first = True
    last_p = None
    for level, text in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = Pt(18)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.0 if level == 1 else 0.6 if level == 2 else 1.2)
        pf.tab_stops.add_tab_stop(Cm(15.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if first:
            # 域开始 + 域指令 + 分隔符
            r = p.add_run()
            fb = OxmlElement('w:fldChar')
            fb.set(qn('w:fldCharType'), 'begin')
            fb.set(qn('w:dirty'), 'true')
            r._r.append(fb)
            r2 = p.add_run()
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = ' TOC \\o "1-3" \\h \\u '
            r2._r.append(it)
            r3 = p.add_run()
            fs = OxmlElement('w:fldChar')
            fs.set(qn('w:fldCharType'), 'separate')
            r3._r.append(fs)
            first = False
        num = (page_map or {}).get(text, 0)
        r = p.add_run('%s\t%s' % (text, num if num else '—'))
        set_run_font(r, ea_font='黑体' if level == 1 else '宋体',
                     size=Pt(12) if level <= 2 else Pt(10.5), bold=(level == 1))
        last_p = p
    if last_p is not None:
        r = last_p.add_run()
        fe = OxmlElement('w:fldChar')
        fe.set(qn('w:fldCharType'), 'end')
        r._r.append(fe)


# --------------------------------------------------------------------------- 主流程
def collect_headings(blocks):
    out = []
    for blk in blocks:
        if blk[0] in ('h1', 'h2', 'h3'):
            out.append((int(blk[0][1]), blk[1]))
    return out


def main():
    page_map = None
    pm_path = os.path.join(ROOT, '_thesis_build', 'tools', 'toc_pages.json')
    if os.path.exists(pm_path):
        import json
        page_map = json.load(open(pm_path, encoding='utf-8'))
        print('using page map with', len(page_map), 'entries')

    doc = Document()
    init_styles(doc)
    init_section(doc.sections[0])

    blocks = load_blocks()
    headings = collect_headings(blocks)

    # 第一段（摘要）之前不加空段；目录之后新建分节，正文页码从 1 开始
    seen_toc = False
    body_section_started = False
    pending_break = False

    for blk in blocks:
        kind = blk[0]
        if kind == 'pagebreak':
            pending_break = True
            continue
        first_para = None
        if kind == 'title':
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(blk[1])
            set_run_font(r, ea_font='黑体', size=Pt(16), bold=True)
            set_outline_level(p, 0)
            first_para = p
        elif kind == 'toctitle':
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(blk[1])
            set_run_font(r, ea_font='黑体', size=Pt(16), bold=True)
            first_para = p
        elif kind == 'toc':
            add_toc(doc, headings, page_map)
            seen_toc = True
        elif kind in ('h1', 'h2', 'h3'):
            level = int(kind[1])
            if seen_toc and not body_section_started and level == 1:
                # 正文另起一节，页码从 1 开始
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                init_section(sec)
                add_page_number_footer(sec, start_at=1)
                body_section_started = True
                pending_break = False
            p = doc.add_heading(blk[1], level=level)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                set_run_font(r, ea_font='黑体')
            first_para = p
        elif kind in ('p', 'p_ind', 'p_kw'):
            first_para = add_body_paragraph(doc, blk[1], indent=(kind == 'p'))
        elif kind == 'ref':
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = Pt(17)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-24)
            r = p.add_run(blk[1])
            set_run_font(r, size=Pt(10.5))
            first_para = p
        elif kind == 'fig':
            first_para = add_figure(doc, blk[1], blk[2])
        elif kind == 'tbl':
            first_para = add_table(doc, blk[1], blk[2], blk[3], blk[4])
        elif kind == 'code':
            first_para = add_code_block(doc, blk[1])
        else:
            raise ValueError('unknown block: ' + kind)
        if pending_break and first_para is not None:
            first_para.paragraph_format.page_break_before = True
            pending_break = False

    doc.save(OUT)
    print('saved', OUT, os.path.getsize(OUT))


if __name__ == '__main__':
    main()
